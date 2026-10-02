from __future__ import annotations
import csv
import io
import json
import os
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path
from xml.etree.ElementTree import ParseError
from .schemas import Dataset,Period
from .financial_import import HEADER_ALIASES,accounting_number

ALIASES={'季度':'period','营业收入':'revenue','营业成本':'cost','净利润':'net_profit','经营现金流':'cash_flow','总资产':'assets','总负债':'liabilities','期初净资产':'equity_begin','期末净资产':'equity_end','库存金额':'inventory','销量':'sales_volume','产量':'production_volume','制造费用':'manufacturing_cost','研发费用':'rd_expense','碳酸锂价格':'lithium_price','行业波动率':'industry_volatility'}

class MalformedWorkbook(ValueError):
    """Recognized XLSX container/XML failures, not unexpected parser defects."""

WORKBOOK_REPAIR='XLSX 文件不完整或内部结构无效；请在表格软件中打开后重新导出为单工作表 XLSX，或改用 CSV。未保存任何导入数据。'

def safe_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>20_000_000:raise ValueError('解压规模超限')
        for i in z.infolist():
            if i.file_size>5_000_000 or '..' in Path(i.filename).parts or i.filename.startswith('/'):raise ValueError('压缩内容不安全')
            if i.compress_size and i.file_size/i.compress_size>1000:raise ValueError('压缩比异常')

def parse_rows(rows,report=None):
    report=report if report is not None else []
    if len(rows)<2:raise ValueError('需要表头和至少一行数据')
    if len(rows)>41 or len(rows[0])>20:raise ValueError('最多40季度、20列')
    aliases={**ALIASES,**HEADER_ALIASES}
    header=[aliases.get(str(v).strip(),str(v).strip()) for v in rows[0]]
    for source,target in zip(rows[0],header):
        if str(source).strip()!=target:report.append({'row':1,'field':target,'source':str(source),'normalized':target,'rule':'表头映射'})
    if len(header)!=len(set(header)):raise ValueError('表头重复')
    if set(header)-set(Period.model_fields):raise ValueError('未知列名：'+','.join(sorted(set(header)-set(Period.model_fields))))
    missing={'period','revenue','cost'}-set(header)
    if missing:raise ValueError('缺少必填列：'+','.join(sorted(missing)))
    result=[]
    for index,row in enumerate(rows[1:],2):
        if not any(v not in ('',None) for v in row):continue
        if len(row)>len(header):raise ValueError(f'第{index}行列数超过表头')
        item={}
        for key,value in zip(header,row):
            if value is None or (isinstance(value,str) and not value.strip()):
                if key in ('period','revenue','cost'):raise ValueError(f'第{index}行{key}不能为空')
                continue
            if isinstance(value,bool):raise ValueError(f'第{index}行不允许布尔数字')
            if key=='period':item[key]=str(value).strip();continue
            try:
                item[key],rule=accounting_number(value,key)
                if rule:report.append({'row':index,'field':key,'source':str(value),'normalized':item[key],'rule':rule})
            except (ValueError,TypeError) as exc:raise ValueError(f'第{index}行 {key}：{exc}') from None
        try:validated=Period.model_validate(item)
        except ValueError as exc:raise ValueError(f'第{index}行校验失败：{exc}') from None
        result.append(validated.model_dump(mode='json'))
    if not result:raise ValueError('没有有效数据行')
    return result

def workbook_rows(data):
    from openpyxl import load_workbook,LXML
    from openpyxl.utils.exceptions import InvalidFileException
    parse_errors=(zipfile.BadZipFile,zlib.error,ParseError,InvalidFileException)
    if LXML:
        from lxml.etree import XMLSyntaxError
        parse_errors+=(XMLSyntaxError,)
    try:
        safe_zip(data)
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if len(wb.sheetnames)!=1:raise ValueError('请只保留一个数据工作表，避免隐式选错表')
            ws=wb.active
            if ws.max_row>41 or ws.max_column>20:raise ValueError('工作表规模超过40季度、20列')
            return [list(row) for row in ws.iter_rows(values_only=True)]
        finally:wb.close()
    except parse_errors:
        raise MalformedWorkbook(WORKBOOK_REPAIR) from None
    except KeyError as exc:
        # ZipFile identifies missing package members this way; do not hide an
        # unrelated programming KeyError behind a user-file validation response.
        if len(exc.args)==1 and isinstance(exc.args[0],str) and exc.args[0].startswith('There is no item named ') and exc.args[0].endswith(' in the archive'):
            raise MalformedWorkbook(WORKBOOK_REPAIR) from None
        raise
    except OSError as exc:
        if str(exc)=='File contains no valid workbook part':
            raise MalformedWorkbook(WORKBOOK_REPAIR) from None
        raise

def import_dataset(filename,data,company,unit='yuan',report=None):
    if len(data)>2_000_000:raise ValueError('文件不得超过2MB')
    ext=Path(filename).suffix.lower()
    if ext=='.json':
        obj=json.loads(data.decode('utf-8-sig'),parse_constant=lambda v:(_ for _ in ()).throw(ValueError('不允许NaN/Infinity')))
        return Dataset.model_validate(obj)
    if ext=='.csv':rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    elif ext=='.xlsx':rows=workbook_rows(data)
    else:raise ValueError('结构化数据只支持JSON、CSV、XLSX；文本证据请使用证据库')
    return Dataset.model_validate({'name':Path(filename).stem,'company':company,'amount_unit':unit,'periods':parse_rows(rows,report)})

def document_text(filename,data):
    if len(data)>2_000_000:raise ValueError('文件不得超过2MB')
    ext=Path(filename).suffix.lower()
    if ext in ('.txt','.md'):text=data.decode('utf-8-sig')
    elif ext=='.pdf':
        if not pdf_status()['enabled']:
            raise ValueError('PDF解析未启用：需要安装requirements-pdf.txt中经单独验收的解析器；当前请上传TXT/MD文本')
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted or len(reader.pages)>60:raise ValueError('不支持加密PDF或超过60页的PDF')
        text='\n'.join((p.extract_text() or '')[:20000] for p in reader.pages)
    else:raise ValueError('证据文件只支持TXT、Markdown、文本型PDF')
    text=text.replace('\x00','').strip()
    if len(text)<20:raise ValueError('文本不足20字符，扫描件请先在可信环境转为文本')
    if len(text)>120000:raise ValueError('提取文本超过120000字符，请拆分文档')
    return text

def parse_document_isolated(filename,data):
    """No untrusted PDF parsing in the long-lived API process; omit secrets from its environment."""
    if len(data)>2_000_000:raise ValueError('文件超过2MB')
    env={k:v for k,v in os.environ.items() if k in ('PATH','SYSTEMROOT','WINDIR','VIRTUAL_ENV','LANG','LC_ALL')};env['PYTHONIOENCODING']='utf-8'
    try:result=subprocess.run([sys.executable,'-m','server.parse_worker',Path(filename).name],input=data,capture_output=True,timeout=12,cwd=Path(__file__).resolve().parent.parent,env=env)
    except subprocess.TimeoutExpired:raise ValueError('文档解析超过12秒，已终止隔离进程') from None
    if len(result.stdout)>800000:raise ValueError('解析输出超限')
    try:obj=json.loads(result.stdout)
    except (ValueError,UnicodeError):raise ValueError('解析进程失败或达到资源限制') from None
    if result.returncode:raise ValueError(obj.get('error','文档解析失败'))
    return obj['text']


def pdf_status():
    from importlib.metadata import version, PackageNotFoundError
    try:
        installed=version('pypdf')
        parts=tuple(int(p) for p in installed.split('.')[:3])
        enabled=parts>=(6,13,2)
    except (PackageNotFoundError,ValueError):
        installed=None;enabled=False
    return {'enabled':enabled,'installed':installed,'minimum':'6.13.2',
        'status':'optional_dependency_available_not_live_verified' if enabled else 'blocked_until_safe_optional_dependency'}
